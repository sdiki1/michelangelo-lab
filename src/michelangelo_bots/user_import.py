from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field, replace
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import select, tuple_
from sqlalchemy.ext.asyncio import AsyncSession

from michelangelo_bots.db import BotUser, datetime_now

MAX_IMPORT_BYTES = 15 * 1024 * 1024
MAX_IMPORT_ROWS = 50_000
MAX_XLSX_UNPACKED_BYTES = 150 * 1024 * 1024
MAX_XLSX_FILES = 5_000
MOSCOW = ZoneInfo("Europe/Moscow")


def normalize_header(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).casefold().replace("ё", "е")
    return re.sub(r"\s+", " ", re.sub(r"[_./\\-]+", " ", normalized)).strip()


@dataclass(frozen=True)
class ImportIssue:
    row: int
    message: str


@dataclass(frozen=True)
class ImportedUser:
    row: int
    platform: str
    platform_user_id: str
    chat_id: str | None = None
    username: str | None = None
    first_name: str | None = None
    last_name: str | None = None
    full_name: str | None = None
    phone: str | None = None
    email: str | None = None
    status: str | None = None
    referral: str | None = None
    comment: str | None = None
    first_seen_at: datetime | None = None
    last_seen_at: datetime | None = None
    total_actions: int = 0

    @property
    def key(self) -> tuple[str, str]:
        return self.platform, self.platform_user_id


@dataclass
class UserImportParseResult:
    filename: str
    total_rows: int = 0
    blank_rows: int = 0
    duplicate_rows: int = 0
    records: list[ImportedUser] = field(default_factory=list)
    errors: list[ImportIssue] = field(default_factory=list)
    warnings: list[ImportIssue] = field(default_factory=list)


@dataclass(frozen=True)
class UserImportApplyResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped_existing: int = 0


FIELD_ALIASES = {
    "platform": {
        "platform",
        "messenger",
        "мессенджер",
        "платформа",
        "источник",
    },
    "platform_user_id": {
        "platform user id",
        "user id",
        "userid",
        "messenger user id",
        "id пользователя",
        "id мессенджера",
        "id пользователя в мессенджере",
    },
    "telegram_user_id": {
        "telegram id",
        "telegram user id",
        "tg id",
        "тг id",
        "id telegram",
        "id телеграм",
    },
    "max_user_id": {
        "max id",
        "max user id",
        "id max",
        "id макс",
    },
    "chat_id": {"chat id", "chatid", "id чата", "чат id"},
    "username": {"username", "user name", "ник", "никнейм", "логин"},
    "first_name": {"first name", "имя"},
    "last_name": {"last name", "surname", "фамилия"},
    "full_name": {"full name", "fio", "фио", "клиент", "имя клиента"},
    "phone": {"phone", "telephone", "телефон", "номер телефона"},
    "email": {"email", "e mail", "почта", "электронная почта"},
    "status": {"status", "статус"},
    "referral": {"referral", "referal", "реферал", "реферер", "источник регистрации"},
    "comment": {"comment", "note", "комментарий", "примечание"},
    "first_seen_at": {
        "first seen at",
        "registered at",
        "registration date",
        "дата регистрации",
        "дата и время регистрации",
    },
    "last_seen_at": {
        "last seen at",
        "last contact",
        "последний контакт",
        "последняя активность",
    },
    "total_actions": {"total actions", "actions", "действий", "количество действий"},
}

ALIAS_TO_FIELD = {
    normalize_header(alias): field_name
    for field_name, aliases in FIELD_ALIASES.items()
    for alias in aliases
}


def parse_user_import(
    filename: str,
    content: bytes,
    *,
    default_platform: str | None = None,
    id_is_platform_user_id: bool = False,
) -> UserImportParseResult:
    safe_name = Path(filename or "users.csv").name
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("Файл больше 15 МБ")
    suffix = Path(safe_name).suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        headers, rows = _read_delimited(content, suffix)
    elif suffix == ".xlsx":
        headers, rows = _read_xlsx(content)
    elif suffix == ".xls":
        raise ValueError("Формат .xls не поддерживается. Сохраните файл как .xlsx или CSV")
    else:
        raise ValueError("Поддерживаются файлы CSV, TSV и XLSX")

    result = UserImportParseResult(filename=safe_name)
    column_map = map_headers(headers, id_is_platform_user_id=id_is_platform_user_id)
    if not column_map:
        raise ValueError("Не найдены знакомые колонки. Скачайте шаблон импорта")

    deduplicated: dict[tuple[str, str], ImportedUser] = {}
    for row_number, raw_row in rows:
        if result.total_rows >= MAX_IMPORT_ROWS:
            raise ValueError(f"В одном файле допускается не более {MAX_IMPORT_ROWS} строк")
        result.total_rows += 1
        if not any(cell_text(value) for value in raw_row):
            result.blank_rows += 1
            continue
        values = {
            field_name: raw_row[index] if index < len(raw_row) else None
            for index, field_name in column_map.items()
        }
        try:
            records, warnings = normalize_row(
                row_number,
                values,
                default_platform=default_platform,
            )
        except ValueError as exc:
            result.errors.append(ImportIssue(row_number, str(exc)))
            continue
        result.warnings.extend(ImportIssue(row_number, warning) for warning in warnings)
        for record in records:
            previous = deduplicated.get(record.key)
            if previous is not None:
                result.duplicate_rows += 1
                record = merge_imported_users(previous, record)
            deduplicated[record.key] = record

    result.records = list(deduplicated.values())
    return result


def map_headers(headers: list[Any], *, id_is_platform_user_id: bool) -> dict[int, str]:
    mapped: dict[int, str] = {}
    used_fields: set[str] = set()
    for index, header in enumerate(headers):
        normalized = normalize_header(cell_text(header))
        field_name = ALIAS_TO_FIELD.get(normalized)
        if normalized == "id" and id_is_platform_user_id:
            field_name = "platform_user_id"
        if field_name and field_name not in used_fields:
            mapped[index] = field_name
            used_fields.add(field_name)
    return mapped


def normalize_row(
    row_number: int,
    values: dict[str, Any],
    *,
    default_platform: str | None,
) -> tuple[list[ImportedUser], list[str]]:
    platform_value = normalize_platform(cell_text(values.get("platform")))
    fallback_platform = normalize_platform(default_platform or "")
    if cell_text(values.get("platform")) and platform_value is None:
        raise ValueError("неизвестный мессенджер; используйте telegram или max")
    if default_platform and fallback_platform is None:
        raise ValueError("неверная платформа по умолчанию")

    telegram_id = identifier_text(values.get("telegram_user_id"))
    max_id = identifier_text(values.get("max_user_id"))
    generic_id = identifier_text(values.get("platform_user_id"))
    chat_id = identifier_text(values.get("chat_id"))
    identities: list[tuple[str, str]] = []
    if telegram_id:
        identities.append(("telegram", telegram_id))
    if max_id:
        identities.append(("max", max_id))
    if generic_id:
        platform = platform_value or fallback_platform
        if platform is None:
            raise ValueError("для User ID не указана платформа")
        identities.append((platform, generic_id))
    if not identities and chat_id:
        platform = platform_value or fallback_platform
        if platform is None:
            raise ValueError("для Chat ID не указана платформа")
        identities.append((platform, chat_id))
    if not identities:
        raise ValueError("нет Platform User ID, Telegram ID или MAX ID")

    first_seen, first_warning = parse_datetime(values.get("first_seen_at"))
    last_seen, last_warning = parse_datetime(values.get("last_seen_at"))
    warnings = [warning for warning in (first_warning, last_warning) if warning]
    username = optional_text(values.get("username"), 255)
    if username:
        username = username.lstrip("@")
    first_name = optional_text(values.get("first_name"), 255)
    last_name = optional_text(values.get("last_name"), 255)
    full_name = optional_text(values.get("full_name"), 512)
    if not full_name:
        full_name = " ".join(part for part in (first_name, last_name) if part) or None
    common = {
        "row": row_number,
        "username": username,
        "first_name": first_name,
        "last_name": last_name,
        "full_name": full_name,
        "phone": optional_text(values.get("phone"), 64),
        "email": normalize_email(values.get("email")),
        "status": normalize_status(values.get("status")),
        "referral": optional_text(values.get("referral"), 255),
        "comment": optional_text(values.get("comment")),
        "first_seen_at": first_seen,
        "last_seen_at": last_seen,
        "total_actions": nonnegative_int(values.get("total_actions")),
    }
    unique_identities = list(dict.fromkeys(identities))
    records = [
        ImportedUser(
            platform=platform,
            platform_user_id=user_id,
            chat_id=(chat_id if len(unique_identities) == 1 else None) or user_id,
            **common,
        )
        for platform, user_id in unique_identities
    ]
    return records, warnings


async def import_user_records(
    session: AsyncSession,
    records: list[ImportedUser],
    *,
    apply: bool,
    existing_mode: str = "merge",
    source_filename: str,
) -> UserImportApplyResult:
    if existing_mode not in {"merge", "overwrite", "skip"}:
        raise ValueError("Неизвестный режим обработки существующих пользователей")
    existing_by_key: dict[tuple[str, str], BotUser] = {}
    keys = [record.key for record in records]
    for chunk in chunks(keys, 500):
        users = (
            await session.execute(
                select(BotUser).where(
                    tuple_(BotUser.platform, BotUser.platform_user_id).in_(chunk)
                )
            )
        ).scalars()
        existing_by_key.update({(user.platform, user.platform_user_id): user for user in users})

    created = updated = unchanged = skipped = 0
    now = datetime_now()
    for record in records:
        user = existing_by_key.get(record.key)
        if user is None:
            created += 1
            if apply:
                first_seen = record.first_seen_at or now
                session.add(
                    BotUser(
                        platform=record.platform,
                        platform_user_id=record.platform_user_id,
                        chat_id=record.chat_id or record.platform_user_id,
                        username=record.username,
                        phone=record.phone,
                        email=record.email,
                        status=record.status or "active",
                        referral=record.referral,
                        comment=record.comment,
                        first_name=record.first_name,
                        last_name=record.last_name,
                        full_name=record.full_name,
                        is_bot=False,
                        raw_profile=legacy_profile(source_filename, record.row),
                        first_seen_at=first_seen,
                        last_seen_at=record.last_seen_at or first_seen,
                        total_actions=record.total_actions,
                        last_action="legacy_import",
                    )
                )
            continue
        if existing_mode == "skip":
            skipped += 1
            continue
        changed = update_existing_user(
            user,
            record,
            overwrite=existing_mode == "overwrite",
            apply=apply,
            source_filename=source_filename,
        )
        if changed:
            updated += 1
        else:
            unchanged += 1

    if apply:
        await session.commit()
    return UserImportApplyResult(created, updated, unchanged, skipped)


def update_existing_user(
    user: BotUser,
    record: ImportedUser,
    *,
    overwrite: bool,
    apply: bool,
    source_filename: str,
) -> bool:
    changed = False
    for attribute in (
        "chat_id",
        "username",
        "phone",
        "email",
        "status",
        "referral",
        "comment",
        "first_name",
        "last_name",
        "full_name",
    ):
        incoming = getattr(record, attribute)
        current = getattr(user, attribute)
        if incoming in (None, "") or (not overwrite and current not in (None, "")):
            continue
        if incoming != current:
            changed = True
            if apply:
                setattr(user, attribute, incoming)
    if record.first_seen_at and record.first_seen_at < user.first_seen_at:
        changed = True
        if apply:
            user.first_seen_at = record.first_seen_at
    if record.last_seen_at and record.last_seen_at > user.last_seen_at:
        changed = True
        if apply:
            user.last_seen_at = record.last_seen_at
    if record.total_actions > (user.total_actions or 0):
        changed = True
        if apply:
            user.total_actions = record.total_actions
    if changed and apply:
        user.raw_profile = {
            **(user.raw_profile or {}),
            **legacy_profile(source_filename, record.row),
        }
    return changed


def merge_imported_users(older: ImportedUser, newer: ImportedUser) -> ImportedUser:
    values: dict[str, Any] = {}
    for attribute in (
        "chat_id",
        "username",
        "first_name",
        "last_name",
        "full_name",
        "phone",
        "email",
        "status",
        "referral",
        "comment",
    ):
        values[attribute] = getattr(newer, attribute) or getattr(older, attribute)
    first_values = [value for value in (older.first_seen_at, newer.first_seen_at) if value]
    last_values = [value for value in (older.last_seen_at, newer.last_seen_at) if value]
    return replace(
        newer,
        **values,
        first_seen_at=min(first_values) if first_values else None,
        last_seen_at=max(last_values) if last_values else None,
        total_actions=max(older.total_actions, newer.total_actions),
    )


def _read_delimited(content: bytes, suffix: str) -> tuple[list[Any], list[tuple[int, list[Any]]]]:
    text = decode_csv(content)
    header_line = 1
    first_line, separator, remainder = text.partition("\n")
    if first_line.rstrip("\r").casefold().startswith("sep="):
        declared = first_line.rstrip("\r")[4:]
        if declared not in {",", ";", "\t", "|"}:
            raise ValueError("Некорректная строка sep= в CSV")
        text = remainder if separator else ""
        delimiter = declared
        header_line = 2
    else:
        delimiter = "\t" if suffix == ".tsv" else None
    sample = text[:8192]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise ValueError("Файл пуст") from exc
    return headers, [
        (index, row) for index, row in enumerate(reader, start=header_line + 1)
    ]


def _read_xlsx(content: bytes) -> tuple[list[Any], list[tuple[int, list[Any]]]]:
    validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось прочитать XLSX-файл") from exc
    try:
        sheet = next((sheet for sheet in workbook.worksheets if sheet.max_row), None)
        if sheet is None:
            raise ValueError("В XLSX нет листов с данными")
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration as exc:
            raise ValueError("Первый лист XLSX пуст") from exc
        rows = [(index, list(row)) for index, row in enumerate(iterator, start=2)]
        return headers, rows
    finally:
        workbook.close()


def validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_FILES:
                raise ValueError("XLSX содержит слишком много внутренних файлов")
            if sum(member.file_size for member in members) > MAX_XLSX_UNPACKED_BYTES:
                raise ValueError("Распакованный XLSX слишком большой")
    except zipfile.BadZipFile as exc:
        raise ValueError("Не удалось прочитать XLSX-файл") from exc


def decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Не удалось определить кодировку CSV; используйте UTF-8 или Windows-1251")


def normalize_platform(value: str) -> str | None:
    normalized = normalize_header(value)
    if normalized in {"telegram", "tg", "телеграм", "тг"}:
        return "telegram"
    if normalized in {"max", "макс"}:
        return "max"
    return None


def normalize_status(value: Any) -> str | None:
    normalized = normalize_header(cell_text(value))
    if not normalized:
        return None
    return {
        "active": "active",
        "активен": "active",
        "активный": "active",
        "lead": "lead",
        "лид": "lead",
        "blocked": "blocked",
        "заблокирован": "blocked",
        "заблокированный": "blocked",
    }.get(normalized, normalized[:64])


def normalize_email(value: Any) -> str | None:
    text = optional_text(value, 255)
    return text.lower() if text else None


def identifier_text(value: Any) -> str | None:
    text = optional_text(value, 128)
    if text and not re.fullmatch(r"[-+A-Za-z0-9_:.]+", text):
        raise ValueError(f"некорректный ID: {text[:40]}")
    return text


def optional_text(value: Any, limit: int | None = None) -> str | None:
    text = cell_text(value)
    if not text:
        return None
    return text[:limit] if limit else text


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse_datetime(value: Any) -> tuple[datetime | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        text = cell_text(value)
        parsed = None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d",
            "%d.%m.%Y %H:%M:%S",
            "%d.%m.%Y %H:%M",
            "%d.%m.%Y",
        ):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            try:
                parsed = datetime.fromisoformat(text)
            except ValueError:
                return None, f"дата «{text[:30]}» не распознана и будет пропущена"
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=MOSCOW)
    return parsed.astimezone(UTC), None


def nonnegative_int(value: Any) -> int:
    try:
        return max(0, int(float(cell_text(value) or "0")))
    except ValueError:
        return 0


def legacy_profile(filename: str, row: int) -> dict[str, Any]:
    return {"legacy_import": {"filename": filename, "row": row}}


def chunks(values: list[Any], size: int) -> Iterable[list[Any]]:
    for index in range(0, len(values), size):
        yield values[index : index + size]
